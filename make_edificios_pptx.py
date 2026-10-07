from pptx import Presentation
from pptx.util import Inches, Pt
import os
BASE = r"C:\Users\diogo\OneDrive\Documentos\Projeto Padrão"
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
def slide(title, bullets):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = title
    tf = sl.placeholders[1].text_frame
    tf.clear()
    for i,b in enumerate(bullets):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        for r in p.runs:
            r.font.size = Pt(15)
    return sl
slide("Montijo — Edifícios: abandonados + novos programas", [
 "Frente ribeirinha Seixalinho–Lançada · base para projeto de arquitetura",
 "Satélite Esri + Nominatim OSM 07/10/2026 · * = aproximado, fixar com GPS",
 "Ficheiros: montijo-website.html · montijo-satelite-edificios.png",
])
slide("Abandonados — essencial", [
 "1 Mundet Bela Vista/Nascentes (1923–88) 38.70080,-8.96927 — naves completas por fora, coberturas/interiores a refazer; medir vão/pé-direito, despistar amianto.",
 "2 SAPAC/ex-ramal 38.70100,-8.96650* — naves em banda, boas para dividir em arrendatários; fixar donos.",
 "3 Rua José Mundet 38.70139,-8.96380 — ~15 000+7 000 m² em ruína parcial; confirmar existência.",
 "4 Armazéns + forno de cal 38.69950,-8.96800* — ruína arqueológica; consolidar, sem mimético.",
 "5 Fábrica completa a leste — grande devoluta 2012; localizar, medir, testar âncora.",
])
slide("Mudaram de programa — referências", [
 "Moinho de Maré do Cais → museu (38.70356,-8.97955): ref.1646, municipal 1995, museu 2005. Lição: maré–caldeira–engenho.",
 "Moinho de Vento do Esteval → museu funcional: 1826, reaberto 2000, moagem ativa; porta SE contra vento.",
 "Cais dos Vapores → lazer/náutica/restauração (38.70302,-8.97687): ponte-cais 1852, edifício 1994, protocolo APL→CM 2025 (467+467+80 m²).",
 "Quinta das Nascentes → Casa da Música Jorge Peixinho: cultura junto ao corredor verde; confirmar programa.",
])
sl = prs.slides.add_slide(prs.slide_layouts[5])
sl.shapes.title.text = "Satélite — 7 edifícios numerados"
img = os.path.join(BASE,"montijo-satelite-edificios.png")
sl.shapes.add_picture(img, Inches(0.5), Inches(1.2), Inches(12.33), Inches(5.9))
slide("Coordenadas + próximos passos", [
 "Verificados: Cais, Moinho do Cais, Bela Vista, J. Mundet, Lançada. Aproximados: SAPAC, forno.",
 "Levantar por edifício: área, vão, pé-direito, estrutura, cobertura, humidade, acesso, dono.",
 "Riscos: amianto, solos, caves em cota de maré. Regra: recuperar e dividir antes de demolir.",
 "Fontes: CM Montijo ARU/ORU/PEZOR · Diário do Distrito 12/06/2025 · ULusófona 2014 · FA-UL 2012.",
])
out = os.path.join(BASE,"montijo-edificios.pptx")
prs.save(out)
print("saved",out)

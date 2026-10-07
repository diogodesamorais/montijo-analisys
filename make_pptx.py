from pptx import Presentation
from pptx.util import Inches, Pt
import os
BASE = r"C:\Users\diogo\OneDrive\Documentos\Projeto Padrão"
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
def add_slide(title, bullets, footer="Montijo — frente ribeirinha abandonada | 2026-10-07 | verificar no local"):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = title
    tf = sl.placeholders[1].text_frame
    tf.clear()
    for i,b in enumerate(bullets):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        for r in p.runs:
            r.font.size = Pt(16)
    tx = sl.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.4)).text_frame
    tx.text = footer
    for p in tx.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
    return sl
add_slide("Montijo — Abandonados da frente ribeirinha", [
 "Pesquisa base para projeto de arquitetura",
 "Aldeia Galega → cidade industrial → frente por reconciliar com o Tejo",
 "Mapa satélite interativo: montijo-mapa-abandonados.html (7 pontos)",
 "ATENÇÃO: propriedade, estrutura, contaminação e cotas exigem visita e ensaios",
])
add_slide("1 · Por que há tanto abandonado — 5 datas", [
 "Até 1930 Aldeia Galega: cais, sal, vinho, pesca, 50 barcos + vapor Lisboa; ~20 matadouros (40 mil porcos/ano).",
 "1908 ramal Pinhal Novo–Montijo: cortiça + carnes junto a rio e estação. Anos 70: ~274 fábricas, 5 358 operários.",
 "1986 UE + 1988–93: carnes fecham (normas), cortiça fecha (plástico/concorrência). Mundet fecha 1988, insolvência 1993. Ramal encerra 1989.",
 "2006–07 requalificação Moinho–Cais sem tratar o miolo industrial.",
 "2025–26: Cais dos Vapores para o município (11 jun 2025); revitalização aprovada por unanimidade (mar 2026).",
])
add_slide("2 · Inventário — o que levantar na visita", [
 "1 Cais dos Vapores 38.70302,-8.97687 — 467+467+80 m², náutica/lúdico/restauração (verificado).",
 "2 Moinho do Cais 38.70356,-8.97955 — maré, museu 2005 (verificado).",
 "3 Bela Vista 38.70080,-8.96927 — primeiros armazéns Mundet 1923 (verificado).",
 "4 Faixa SAPAC 38.70100,-8.96650 — naves em banda (APROX., fixar GPS).",
 "5 Rua José Mundet 38.70139,-8.96380 — ~15 000+7 000 m² em ruína (verificar existência).",
 "6 Forno de cal 38.69950,-8.96800 — ruína arqueológica (APROX.).",
 "7 Leste + Lançada 38.68875,-8.95138 — fábrica completa (verificar) + moinho 1368.",
])
add_slide("3 · O que o plano permite — e obriga", [
 "PEZOR 267 ha (Seixalinho–Lançada): serviços/lazer na urbana, reavaliar nascente, natural só recreio + ciclável.",
 "Recentrar 157 ha: consolidar, conter norte, aproximar do rio; corredor ecológico.",
 "ARU 274 ha: núcleo prioritário + marginal nascente + corredor Assentas–Esteval; devolutos = potenciais benefícios.",
 "PDM 1997 em revisão (5 eixos 2022). Condicionantes: APL/APA, privados fragmentados, maré/cheia, rail.",
])
add_slide("4 · Estratégia de projeto — 5 movimentos", [
 "1 Levantar antes de desenhar: matriz área/vão/pé-direito/estrutura/cobertura/humidade/dono.",
 "2 Não demolir por defeito: dividir em arrendatários menores; criativas + serviços + habitação pontual.",
 "3 Coser rio–ferrovia–cidade: ramal como alameda/ciclovia; R/C permeável à maré, programa em cima.",
 "4 Energia do sítio: inércia + cruzada/noturna + sombra S/SO + cisterna; despistar amianto.",
 "5 Fasear pelo Cais: âncora náutica/restauração e estender a nascente por quarteirões-piloto.",
])
sl = add_slide("5 · Mapa satélite", [
 "Abrir no browser: montijo-mapa-abandonados.html (Esri World Imagery, 7 marcadores).",
 "Linha amarela: faixa rio–ferrovia–fábrica.",
 "Verde = Nominatim/CM verificado; amarelo = aprox. a fixar com GPS.",
 "Prints de apoio: montijo-centro-mapa.png e montijo-mapa-arvores.png na pasta.",
])
# try embed static maps if present
for img in ["montijo-centro-mapa.png","montijo-mapa-arvores.png"]:
    p = os.path.join(BASE, img)
    if os.path.exists(p):
        try:
            sl.shapes.add_picture(p, Inches(7.5), Inches(1.4), Inches(5.0), Inches(3.6))
            break
        except Exception as e:
            print("img skip", e)
add_slide("Fontes + próximos passos", [
 "CM Montijo (património, ARU/ORU, PEZOR, drenagem) · Diário do Distrito 12/06/2025 · O Setubalense 28/10/2022.",
 "ULusófona 2014 · FA-UL 2012 · estágio CM Montijo 2012 · ESCS 2022 · Nominatim OSM 07/10/2026.",
 "Próximos: certidões, GPS, estrutura/cobertura, amianto/solos, cotas maré/cheia, conversa CM/APL.",
])
out = os.path.join(BASE, "montijo-frente-ribeirinha.pptx")
prs.save(out)
print("saved", out)

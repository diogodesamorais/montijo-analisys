from pptx import Presentation
from pptx.util import Inches, Pt
import os
BASE = r"C:\Users\diogo\OneDrive\Documentos\Projeto Padrão"
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def slide(title, bullets, fsize=15, footer="Montijo ribeirinha — website 07/10/2026 | Esri + OSM | verificar no local"):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = title
    tf = sl.placeholders[1].text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        for r in p.runs:
            r.font.size = Pt(fsize)
    tx = sl.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.4)).text_frame
    tx.text = footer
    for p in tx.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
    return sl

def imgslide(title, imgfile, caption=""):
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = title
    p = os.path.join(BASE, imgfile)
    sl.shapes.add_picture(p, Inches(0.5), Inches(1.2), Inches(12.33), Inches(5.5))
    if caption:
        tx = sl.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.4)).text_frame
        tx.text = caption
        for par in tx.paragraphs:
            for r in par.runs:
                r.font.size = Pt(10)
    return sl

slide("Montijo ribeirinha — do abandono à reconversão", [
 "Website de projeto: index.html — 8 abandonados (1–8) + 8 mudanças de função (M1–M8).",
 "Faixa rio–ferrovia–fábrica sobre satélite Esri World Imagery, z16, 07/10/2026.",
 "Fichas: montijo-ribeirinha-abandonados-satelite.html + montijo-ribeirinha-mudanca-funcao.html.",
 "ATENÇÃO: pontos aprox. a fixar com GPS, certidões, ARU/PEZOR e cotas de maré.",
], fsize=16)

slide("1 · Abandonados 1–4 (ver tabela no site §1)", [
 "1 Cais dos Vapores 38.70302,-8.97687 — terminal em transição; futuro lazer/náutica/restauração.",
 "2 Mundet Bela Vista 1923 38.70080,-8.96927 — naves cortiça; estrutura completa, coberturas a refazer.",
 "3 Chaminé + Caldeiras* 38.70050,-8.96980 — exposição desde 2025, restante devoluto.",
 "4 SAPAC / ex-ramal* 38.70100,-8.96650 — naves em banda; confirmar troço a troço (* = aprox.).",
])

slide("2 · Abandonados 5–8", [
 "5 Rua José Mundet 38.70139,-8.96380 — ruína parcial ~15 000 + 7 000 m²; confirmar cadastro.",
 "6 Forno de cal* 38.69950,-8.96800 — ruína arqueológica em cota baixa; risco maré.",
 "7 Antiga Estação + ramal 1908–89* 38.70230,-8.97180 — levantar restos; troços em ecopista.",
 "8 Leste + Moinho Lançada 1368 38.68875,-8.95138 — fábrica a confirmar + moinho antigo.",
])

imgslide("3 · Satélite — 8 abandonados numerados", "montijo-satelite-abandonados-v2.png",
 "Esri World Imagery z16 | nºs = tabela §1 do site | * fixar com GPS")

slide("4 · Mudaram de função M1–M4", [
 "M1 Moinho do Cais → museu: ref. 1646, municipal 1995, museu 2005–06 (38.70356,-8.97955).",
 "M2 Casa Mora → museu: casa ~1875, biblioteca 80s, museu 1993 (38.70667,-8.97617).",
 "M3 Mercado → reabilitado: âncora comercial (38.70621,-8.97624).",
 "M4 Touros → Cinema-Teatro Joaquim d'Almeida 1957 (38.70588,-8.97037).",
])

slide("5 · Mudaram de função M5–M8", [
 "M5 Cais → lazer/náutica/restauração (em curso): protocolo APL→CM 11/06/2025.",
 "M6 Armazéns R. Miguel Pais → passeio 2006–07 (38.70381,-8.97563).",
 "M7 Ramal 1908–89 → Ecopista Montijo Ciclável* (38.70230,-8.97180).",
 "M8 Caldeiras Mundet → exposição cortiça desde 2025* (38.70050,-8.96980).",
])

imgslide("6 · Satélite — 8 mudanças de função", "montijo-satelite-mudanca-funcao.png",
 "Esri z16 | M1–M8 = tabela §2 do site | * aprox. a confirmar")

slide("7 · Cronologia + estratégia de projeto", [
 "Até 1930 Aldeia Galega rio-economia → 1908 ramal/indústria → 1986–93 fechos → 2006–07 passeio → 2025–26 janela Cais.",
 "1 Levantar antes de desenhar (vão/pé-direito/estrutura/cota/dono). 2 Dividir naves, não demolir.",
 "3 Ex-ramal como alameda/ciclável; R/C permeável. 4 Inércia + cruzada + sombra + cisterna; amianto.",
 "5 Fasear pelo Cais a nascente por quarteirões-piloto (nºs 2–5).",
])

slide("Fontes + próximos passos", [
 "CM Montijo (Moinho, Casa Mora, Cinema, ARU/ORU, PEZOR, Ciclável, Mercado) · protocolo APL–CM 11/06/2025.",
 "AMRS Caldeiras 2025 · tese ULisboa Folgado · ULusófona 2014 · FA-UL 2012 · Nominatim 07/10/2026.",
 "Abrir: index.html (site) · fichas html · PNG satélite. Pedir certidões + PDFs ARU/PEZOR; ir com GPS.",
])

out = os.path.join(BASE, "montijo-ribeirinha-website.pptx")
prs.save(out)
print("saved", out)
